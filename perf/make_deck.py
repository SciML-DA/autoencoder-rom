"""
make_deck.py
============

Build the weekly presentation from whatever plots currently exist on disk.

Regenerating beats hand-editing: the plots get rebuilt every time a sweep is
rerun, and a deck assembled by hand goes stale silently. Slides are only added
for figures that are actually present, so this can be run before the sweeps
finish and again afterwards.

    python perf/make_deck.py --out "/path/to/PresentationWk6.pptx"
"""

from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

W, H = Inches(13.333), Inches(7.5)
DEEP = RGBColor(0x06, 0x5A, 0x82)
TEAL = RGBColor(0x1C, 0x72, 0x93)
MID = RGBColor(0x21, 0x29, 0x5C)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

CONV = Path("results/convergence_new/bl")
CONV_OLD = Path("results/convergence/bl")
PERF = Path("results/performance/baseline")
SWEEP = Path("results/sweeps")


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def text(slide, s, x, y, w, h, size=16, bold=False, color=INK,
         align=PP_ALIGN.LEFT, font="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    for i, line in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return tb


def fit(slide, img: Path, x, y, box_w, box_h):
    """Place an image centred in a box, preserving aspect ratio."""
    iw, ih = Image.open(img).size
    scale = min(box_w / iw, box_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(str(img), Emu(int(x + (box_w - w) / 2)),
                             Emu(int(y + (box_h - h) / 2)), Emu(w), Emu(h))


def title_slide(prs, title, subtitle):
    s = blank(prs)
    bg(s, MID)
    text(s, title, Inches(0.9), Inches(2.5), Inches(11.5), Inches(1.4),
         size=44, bold=True, color=WHITE, font="Cambria")
    text(s, subtitle, Inches(0.9), Inches(4.0), Inches(11.5), Inches(1.2),
         size=18, color=RGBColor(0xCA, 0xDC, 0xFC))
    return s


def section(prs, kicker, title):
    s = blank(prs)
    bg(s, WHITE)
    text(s, kicker.upper(), Inches(0.7), Inches(0.5), Inches(11.9), Inches(0.3),
         size=12, bold=True, color=TEAL)
    text(s, title, Inches(0.7), Inches(0.85), Inches(11.9), Inches(0.8),
         size=30, bold=True, color=INK, font="Cambria")
    return s


def figure_slide(prs, kicker, title, imgs, caption=None):
    imgs = [p for p in imgs if p.exists()]
    if not imgs:
        return None
    s = section(prs, kicker, title)
    top = Inches(1.85)
    avail_h = Inches(4.9) if caption else Inches(5.3)
    if len(imgs) == 1:
        fit(s, imgs[0], Inches(0.7), top, Inches(11.9), avail_h)
    else:
        gap = Inches(0.25)
        each = int((Inches(11.9) - gap * (len(imgs) - 1)) / len(imgs))
        for i, p in enumerate(imgs):
            fit(s, p, Inches(0.7) + i * (each + gap), top, each, avail_h)
    if caption:
        text(s, caption, Inches(0.7), Inches(6.85), Inches(11.9), Inches(0.5),
             size=12, color=MUTED)
    return s


def build(out: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    title_slide(prs, "Autoencoder ROM — Week 6",
                "Convergence study on the boundary-layer dataset, "
                "GPU performance baseline, and JAX hyperparameter sweeps")

    # --- what changed -----------------------------------------------------
    s = section(prs, "Summary", "What I did this week")
    bullets = [
        ("Cluster", "Moved the study onto the Aero GPU cluster (meteor, 3x T4). "
                    "Pre-decimated the 35 GB dataset to 4.4 GB with identical "
                    "loaded arrays, so transfers and quota stopped being a problem."),
        ("JAX on GPU", "Found JAX was running CPU-only — the CUDA plugin was never "
                       "installed. Fixed; JAX and torch now agree to within 5% on a "
                       "matched matmul instead of JAX looking 90x slower."),
        ("Performance", "Built a timing baseline over 192 configurations and "
                        "optimised the training loops: AE +33%, AEJax +50%."),
        ("Sweeps", "Eight-point sweeps over latent size, learning rate, weight decay, "
                   "batch size, width and input noise, JAX models only."),
    ]
    y = Inches(1.9)
    for head, body in bullets:
        text(s, head, Inches(0.7), y, Inches(2.2), Inches(0.4), size=15, bold=True, color=DEEP)
        text(s, body, Inches(3.0), y, Inches(9.6), Inches(1.0), size=13, color=INK)
        y += Inches(1.15)

    # --- convergence study ------------------------------------------------
    figure_slide(prs, "Convergence study — bl", "Model size and fit cost",
                 [CONV / "cost.png"],
                 "Re-run 12 Aug on GPU. The 7 Aug figures had JAX on CPU, inflating AEJax at "
                 "k=256 by 47x (9158 s vs 196 s). Accuracy was unaffected — it reproduced to "
                 "within 1%.")
    figure_slide(prs, "Convergence study — bl", "Accuracy and generalisation vs latent size",
                 [CONV / "generalisation.png"],
                 "Right panel: test error divided by train error. AE/AEJax reach a 20x gap at "
                 "k=256 — they reconstruct unseen snapshots 20x worse than fitted ones, while "
                 "POD and the conv models stay near 1.")

    curves = sorted(CONV_OLD.glob("loss_curves_latent*.png"),
                    key=lambda p: int(p.stem.rsplit("latent", 1)[1]))
    for i in range(0, len(curves), 4):
        chunk = curves[i:i + 4]
        ks = ", ".join(p.stem.rsplit("latent", 1)[1] for p in chunk)
        figure_slide(prs, "Convergence study — bl", f"Training curves, latent {ks}", chunk)

    # --- performance ------------------------------------------------------
    figure_slide(prs, "Performance baseline", "What actually drives compute cost",
                 [PERF / "time_vs_params.png"],
                 "192 configs on 3x T4, one per GPU. Dense AE is flat below ~7.2e6 parameters "
                 "then turns compute-bound; conv spread is architecture, not noise.")
    figure_slide(prs, "Performance baseline", "Latent size is nearly free",
                 [PERF / "time_vs_latent.png"],
                 "The flat lines are fixed-width variants: at fixed width, a 128x increase in "
                 "latent changes runtime by under 3%. The study's default ties width to latent.")
    figure_slide(prs, "Performance baseline", "Compile cost and JAX vs torch",
                 [PERF / "compile_cost.png", PERF / "jax_ratio.png"],
                 "JAX pays 5-7x more one-time compile. Its speed advantage is dataset-dependent: "
                 "conv-JAX is 1.4x faster on bl but 4x slower on circle's larger grid.")

    # --- optimisation results --------------------------------------------
    s = section(prs, "Optimisation", "Training-loop changes, measured")
    rows = [
        ("Fused Adam", "AE +26% (k=64), +33% (k=256)", "Adam was 76% of CUDA time"),
        ("Gather inside jit", "AEJax +50% (k=8)", "bitwise identical, 0.00e+00"),
        ("Snapshot buffers", "+1-2%", "removed per-improvement deepcopy"),
        ("torch.compile", "rejected — 0.97x", "the 1.96x probe had changed the workload"),
        ("NHWC layout", "rejected — 1.01x", "XLA already handles layout"),
        ("lax.scan / vmap", "rejected", "vmap 1.46x at k=8 but 0.80x at k=64"),
    ]
    y = Inches(1.95)
    for a, b, c in rows:
        text(s, a, Inches(0.7), y, Inches(3.0), Inches(0.4), size=14, bold=True, color=DEEP)
        text(s, b, Inches(3.8), y, Inches(3.4), Inches(0.4), size=14, color=INK)
        text(s, c, Inches(7.4), y, Inches(5.2), Inches(0.4), size=13, color=MUTED)
        y += Inches(0.62)
    text(s, "A correctness bug also surfaced: a non-blocking index transfer let training run on "
            "partially-written batches. Invisible on CUDA, caught on MPS.",
         Inches(0.7), Inches(6.1), Inches(11.9), Inches(0.8), size=13, color=INK)

    # --- sweeps (only if they have finished) ------------------------------
    sweep_pngs = sorted(SWEEP.glob("*.png"))
    for p in sweep_pngs:
        figure_slide(prs, "Hyperparameter sweeps — bl (JAX)",
                     p.stem.replace("_", " "), [p])

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"wrote {out}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    build(Path(ap.parse_args().out).expanduser())
